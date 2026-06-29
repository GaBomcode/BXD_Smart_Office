"""AI Draft Engine service.

The engine is advisory only: it prepares templates, evidence, outlines, and
draft text for human review. It does not publish or export before approval.
"""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
import json
import logging
from pathlib import Path
import re
import unicodedata
from typing import Any

from core.config import DOCUMENT_EXPORT_DIR
from models.ai_draft_citation import AIDraftCitation
from models.ai_draft_request import AIDraftRequest
from models.ai_draft_result import AIDraftResult
from repositories.ai_draft_repository import AIDraftRepository
from repositories.document_library_repository import DocumentLibraryRepository
from services.document_validation.draft_validator import DraftValidator
from services.document_service import DocumentService
from services.knowledge_service import KnowledgeService

logger = logging.getLogger(__name__)

SUPPORTED_DOCUMENT_TYPES = ["Cong van", "Bao cao", "Ke hoach", "To trinh", "Thong bao", "Giay moi"]
LOW_CONFIDENCE_THRESHOLD = 0.55


@dataclass(slots=True)
class IntentDetection:
    document_type: str | None
    confidence: float
    requires_user_selection: bool
    matched_signals: list[str]

    def to_dict(self) -> dict[str, Any]:
        return {
            "document_type": self.document_type,
            "confidence": self.confidence,
            "requires_user_selection": self.requires_user_selection,
            "matched_signals": self.matched_signals,
        }


class AIDraftService:
    """Coordinates AI Draft workflow with human review gates."""

    def __init__(
        self,
        *,
        repository: AIDraftRepository | None = None,
        library_repository: DocumentLibraryRepository | None = None,
        knowledge_service: KnowledgeService | None = None,
        document_service: DocumentService | None = None,
    ) -> None:
        self.repository = repository or AIDraftRepository()
        self.library_repository = library_repository or DocumentLibraryRepository()
        self.knowledge_service = knowledge_service or KnowledgeService(library_repository=self.library_repository)
        self.document_service = document_service or DocumentService()

    def create_request(self, request_text: str, *, requested_by: str = "Người dùng") -> int:
        """Create an AI Draft request after rule-based intent detection."""
        detection = self.detect_intent(request_text)
        status = "needs_document_type_selection" if detection.requires_user_selection else "ready_for_template_selection"
        request = AIDraftRequest(
            request_text=request_text,
            requested_by=requested_by,
            detected_document_type=detection.document_type,
            selected_document_type=detection.document_type if not detection.requires_user_selection else None,
            confidence=detection.confidence,
            status=status,
        )
        request_id = self.repository.create_request(request)
        self.repository.log_audit(
            action="create_ai_draft_request",
            entity_type="ai_draft_request",
            entity_id=request_id,
            actor=requested_by,
            detail=json.dumps(detection.to_dict(), ensure_ascii=False),
        )
        return request_id

    def detect_intent(self, request_text: str) -> IntentDetection:
        """Detect document type with deterministic rules, no LLM."""
        normalized = self._normalize(request_text)
        if not normalized:
            raise ValueError("Yêu cầu soạn thảo không được rỗng")
        rules: dict[str, list[str]] = {
            "Cong van": ["cong van", "de nghi", "kinh gui", "phoi hop", "trien khai"],
            "Bao cao": ["bao cao", "ket qua", "tinh hinh", "danh gia", "kien nghi"],
            "Ke hoach": ["ke hoach", "muc dich", "yeu cau", "tien do", "to chuc thuc hien"],
            "To trinh": ["to trinh", "trinh", "xin chu truong", "phe duyet", "xem xet"],
            "Thong bao": ["thong bao", "thong tin", "ket luan", "lich lam viec"],
            "Giay moi": ["giay moi", "moi hop", "tham du", "hoi nghi", "cuoc hop"],
        }
        scored: list[tuple[str, int, list[str]]] = []
        for document_type, signals in rules.items():
            matched = [signal for signal in signals if signal in normalized]
            score = len(matched)
            if self._normalize(document_type) in normalized:
                score += 2
            scored.append((document_type, score, matched))
        scored.sort(key=lambda item: item[1], reverse=True)
        best_type, best_score, matched = scored[0]
        if best_score == 0:
            return IntentDetection(None, 0.0, True, [])
        second_score = scored[1][1] if len(scored) > 1 else 0
        confidence = min(0.35 + best_score * 0.18 + max(best_score - second_score, 0) * 0.08, 0.98)
        confidence = round(confidence, 2)
        return IntentDetection(best_type, confidence, confidence < LOW_CONFIDENCE_THRESHOLD, matched)

    def select_document_type(self, request_id: int, document_type: str, *, actor: str = "Người dùng") -> dict[str, Any]:
        """Human selects document type when rule confidence is low."""
        if document_type not in SUPPORTED_DOCUMENT_TYPES:
            raise ValueError("Loại văn bản không hợp lệ")
        self._require_request(request_id)
        self.repository.update_request(
            request_id,
            {
                "selected_document_type": document_type,
                "status": "ready_for_template_selection",
                "confidence": 1.0,
            },
        )
        self.repository.log_audit(
            action="select_ai_draft_document_type",
            entity_type="ai_draft_request",
            entity_id=request_id,
            actor=actor,
            detail=document_type,
        )
        return self._require_request(request_id)

    def suggest_templates(self, request_id: int, *, limit: int = 5) -> list[dict[str, Any]]:
        """Return nearest internal document templates with citations."""
        request = self._require_request(request_id)
        document_type = self._selected_type(request)
        request_words = set(self._tokens(str(request["request_text"])))
        candidates = self.library_repository.list_documents(status=None, document_type=None, limit=None)
        suggestions: list[dict[str, Any]] = []
        for document in candidates:
            score = 0.0
            reasons: list[str] = []
            if self._same_type(document.get("document_type"), document_type):
                score += 0.35
                reasons.append("Cung loai van ban")
            if self._is_internal_agency(document.get("issuing_agency")):
                score += 0.25
                reasons.append("Uu tien van ban noi bo Ban Xay dung Dang")
            text = " ".join(str(document.get(key) or "") for key in ("title", "summary", "keywords", "document_number"))
            matched = request_words.intersection(self._tokens(text))
            if matched:
                score += min(len(matched) * 0.06, 0.3)
                reasons.append("Trung tu khoa: " + ", ".join(sorted(matched)[:5]))
            if str(document.get("status") or "") in {"indexed", "need_review"}:
                score += 0.1
            if score <= 0:
                continue
            citation = self._citation_from_library_document(document, score=score, reason="; ".join(reasons))
            suggestions.append(
                {
                    "document": document,
                    "score": round(min(score, 1.0), 4),
                    "reason": "; ".join(reasons) if reasons else "Mẫu trong kho văn bản",
                    "citation": citation.to_dict(),
                }
            )
        top = sorted(suggestions, key=lambda item: item["score"], reverse=True)[:limit]
        self.repository.replace_citations(
            request_id=request_id,
            result_id=None,
            citation_type="template",
            citations=[
                self._citation_from_library_document(
                    item["document"],
                    score=float(item["score"]),
                    reason=str(item["reason"]),
                )
                for item in top
            ],
        )
        return top

    def collect_evidence(self, request_id: int, *, top_k: int = 5) -> list[dict[str, Any]]:
        """Collect source evidence through Knowledge Engine semantic search."""
        request = self._require_request(request_id)
        query = str(request["request_text"])
        try:
            results = self.knowledge_service.semantic_search(query, top_k=top_k)
        except Exception as exc:
            logger.warning("Semantic search failed for AI draft request %s: %s", request_id, exc)
            results = []
        citations = [self._citation_from_search_result(result) for result in results]
        self.repository.replace_citations(
            request_id=request_id,
            result_id=None,
            citation_type="evidence",
            citations=citations,
        )
        return [
            {
                "document": result.get("document"),
                "chunk": result.get("chunk"),
                "score": result.get("score"),
                "citation": citation.to_dict(),
            }
            for result, citation in zip(results, citations, strict=False)
        ]

    def generate_outline(
        self,
        request_id: int,
        *,
        template_document_id: int | None = None,
        title: str | None = None,
    ) -> int:
        """Generate an outline and persist it for human review."""
        request = self._require_request(request_id)
        document_type = self._selected_type(request)
        evidence = self.repository.list_citations(request_id=request_id, citation_type="evidence")
        template = self.library_repository.get_document(template_document_id) if template_document_id else None
        outline = self._compose_outline(
            request_text=str(request["request_text"]),
            document_type=document_type,
            template=template,
            evidence=evidence,
        )
        result = AIDraftResult(
            request_id=request_id,
            title=title or f"Dự thảo {document_type.lower()}",
            document_type=document_type,
            template_document_id=template_document_id,
            outline_content=outline,
            status="pending_outline_review",
        )
        result_id = self.repository.create_result(result)
        for citation_type in ("template", "evidence"):
            existing = self.repository.list_citations(request_id=request_id, citation_type=citation_type)
            copied = [self._citation_from_row(row) for row in existing]
            self.repository.replace_citations(
                request_id=request_id,
                result_id=result_id,
                citation_type=citation_type,
                citations=copied,
            )
        self.repository.update_request(request_id, {"status": "pending_outline_review"})
        return result_id

    def approve_outline(self, result_id: int, *, approved_outline: str, actor: str = "Người dùng") -> dict[str, Any]:
        """Human approves or edits outline before draft generation."""
        result = self._require_result(result_id)
        if str(result["status"]) != "pending_outline_review":
            raise ValueError("Dàn ý không ở trạng thái chờ duyệt")
        if not approved_outline.strip():
            raise ValueError("Dàn ý đã duyệt không được rỗng")
        self.repository.update_result(
            result_id,
            {"outline_content": approved_outline.strip(), "status": "outline_approved", "review_note": "Outline approved"},
        )
        self.repository.add_revision(
            result_id=result_id,
            content=approved_outline.strip(),
            status="outline_approved",
            edited_by=actor,
            note="Human approved outline",
        )
        self.repository.log_audit(
            action="approve_ai_draft_outline",
            entity_type="ai_draft_result",
            entity_id=result_id,
            actor=actor,
        )
        return self._require_result(result_id)

    def generate_draft(self, result_id: int) -> dict[str, Any]:
        """Generate full draft only after outline approval."""
        result = self._require_result(result_id)
        if str(result["status"]) != "outline_approved":
            raise ValueError("Cần duyệt dàn ý trước khi sinh nội dung")
        request = self._require_request(int(result["request_id"]))
        citations = self.repository.list_citations(result_id=result_id)
        draft = self._compose_draft(
            request_text=str(request["request_text"]),
            result=result,
            citations=citations,
        )
        self.repository.update_result(result_id, {"draft_content": draft, "status": "pending_user_review"})
        self.repository.add_revision(
            result_id=result_id,
            content=draft,
            status="pending_user_review",
            edited_by="AI Draft Engine",
            note="Rule-based draft generated with citations",
        )
        return self._require_result(result_id)

    def save_user_review(
        self,
        result_id: int,
        edited_content: str,
        *,
        actor: str = "Người dùng",
        note: str | None = None,
    ) -> dict[str, Any]:
        """Save user edits without publishing the draft."""
        result = self._require_result(result_id)
        if str(result["status"]) not in {"pending_user_review", "approved"}:
            raise ValueError("Dự thảo chưa sẵn sàng để chỉnh sửa")
        if not edited_content.strip():
            raise ValueError("Nội dung dự thảo không được rỗng")
        status = "pending_user_review"
        self.repository.update_result(result_id, {"draft_content": edited_content.strip(), "status": status, "review_note": note})
        self.repository.add_revision(
            result_id=result_id,
            content=edited_content.strip(),
            status=status,
            edited_by=actor,
            note=note,
        )
        return self._require_result(result_id)

    def approve_draft(self, result_id: int, *, actor: str = "Người dùng", note: str | None = None) -> dict[str, Any]:
        """Human approval gate before export."""
        result = self._require_result(result_id)
        if str(result["status"]) != "pending_user_review":
            raise ValueError("Dự thảo phải ở trạng thái chờ người dùng duyệt")
        if not str(result.get("draft_content") or "").strip():
            raise ValueError("Dự thảo chưa có nội dung")
        now = datetime.now().isoformat(timespec="seconds")
        self.repository.update_result(
            result_id,
            {"status": "approved", "approved_by": actor, "approved_at": now, "review_note": note},
        )
        self.repository.log_audit(
            action="approve_ai_draft",
            entity_type="ai_draft_result",
            entity_id=result_id,
            actor=actor,
            detail=note,
        )
        return self._require_result(result_id)

    def export_docx(self, result_id: int) -> Path:
        """Export only approved drafts to DOCX."""
        result = self._require_result(result_id)
        if str(result["status"]) != "approved":
            raise ValueError("Chỉ xuất DOCX sau khi người dùng đã duyệt")
        content = str(result.get("draft_content") or "").strip()
        if not content:
            raise ValueError("Dự thảo chưa có nội dung")
        validation = DraftValidator(
            repository=self.repository,
            library_repository=self.library_repository,
        ).validate(result)
        if not validation.passed:
            raise ValueError("Draft validation failed: " + "; ".join(validation.errors))
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm, Pt
        except ImportError as exc:
            raise RuntimeError("Thiếu python-docx để xuất DOCX") from exc

        settings = self.document_service.get_document_settings()
        DOCUMENT_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        output = DOCUMENT_EXPORT_DIR / f"ai_draft_{result_id:04d}_{self._slugify(str(result['title']))}.docx"
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(float(settings["margin_top_cm"]))
        section.bottom_margin = Cm(float(settings["margin_bottom_cm"]))
        section.left_margin = Cm(float(settings["margin_left_cm"]))
        section.right_margin = Cm(float(settings["margin_right_cm"]))
        style = document.styles["Normal"]
        style.font.name = settings["font_name"]
        style.font.size = Pt(float(settings["font_size"]))
        for line in content.splitlines():
            paragraph = document.add_paragraph()
            text = line.strip()
            if not text:
                continue
            run = paragraph.add_run(text)
            if text.isupper() or text.startswith(("AI DRAFT", "DU THAO")):
                run.bold = True
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        document.save(output)
        self.repository.update_result(result_id, {"output_path": str(output), "status": "exported"})
        return output

    def get_result_package(self, result_id: int) -> dict[str, Any]:
        """Return result with citations and revisions for UI rendering."""
        result = self._require_result(result_id)
        return {
            "result": result,
            "request": self._require_request(int(result["request_id"])),
            "citations": self.repository.list_citations(result_id=result_id),
            "revisions": self.repository.list_revisions(result_id),
        }

    def list_requests(self) -> list[dict[str, Any]]:
        return self.repository.list_requests(limit=100)

    def list_results(self) -> list[dict[str, Any]]:
        return self.repository.list_results(limit=100)

    def _compose_outline(
        self,
        *,
        request_text: str,
        document_type: str,
        template: dict[str, Any] | None,
        evidence: list[dict[str, Any]],
    ) -> str:
        source_note = "Có nguồn căn cứ kèm citation." if evidence else "Chưa có nguồn căn cứ; không được khẳng định chắc chắn."
        template_note = f"Mẫu tham chiếu: {template.get('title') or template.get('file_name')}" if template else "Chưa chọn mẫu."
        sections = self._outline_sections(document_type)
        lines = [
            f"Dàn ý {document_type}",
            f"Yêu cầu: {request_text.strip()}",
            template_note,
            source_note,
            "",
        ]
        for index, section in enumerate(sections, start=1):
            lines.append(f"{index}. {section}")
            if evidence:
                lines.append("   - Gắn citation liên quan trước khi viết nội dung.")
            else:
                lines.append("   - Chỉ nêu dự kiến tham mưu, chờ người dùng bổ sung căn cứ.")
        return "\n".join(lines)

    def _compose_draft(self, *, request_text: str, result: dict[str, Any], citations: list[dict[str, Any]]) -> str:
        document_type = str(result["document_type"])
        citation_lines = self._citation_lines(citations)
        if citation_lines:
            evidence_intro = "Căn cứ các nguồn đã thu thập:"
            evidence_block = "\n".join(citation_lines)
        else:
            evidence_intro = "Chưa có nguồn căn cứ được tìm thấy."
            evidence_block = "Nội dung dưới đây chỉ là gợi ý tham mưu, cần người dùng bổ sung/kiểm chứng căn cứ."
        outline = str(result.get("outline_content") or "")
        settings = self.document_service.get_document_settings()
        return "\n".join(
            [
                "DỰ THẢO",
                str(result["title"]).upper(),
                "",
                settings["agency_name"],
                f"Số: ......../{settings['document_code_prefix']}",
                "",
                f"Loại văn bản: {document_type}",
                f"Yêu cầu người dùng: {request_text.strip()}",
                "",
                evidence_intro,
                evidence_block,
                "",
                "Nội dung tham mưu:",
                self._draft_body_from_outline(outline, has_evidence=bool(citation_lines)),
                "",
                "Lưu ý: AI Draft Engine chỉ tham mưu. Người dùng phải kiểm tra, chỉnh sửa và duyệt trước khi xuất/ban hành.",
                "",
                "Nơi nhận:",
                "- Như trên;",
                "- Lưu: VT.",
                "",
                settings["signer_title"],
                settings["signer_name"],
            ]
        )

    def _draft_body_from_outline(self, outline: str, *, has_evidence: bool) -> str:
        section_lines = [line for line in outline.splitlines() if re.match(r"^\d+\.", line.strip())]
        body: list[str] = []
        for line in section_lines:
            title = re.sub(r"^\d+\.\s*", "", line.strip())
            body.append(line.strip())
            if has_evidence:
                body.append(f"Trình bày nội dung về {title.lower()} theo các citation đã gắn, không mở rộng ngoài nguồn.")
            else:
                body.append(f"Đề xuất người dùng bổ sung căn cứ trước khi khẳng định về {title.lower()}.")
            body.append("")
        return "\n".join(body).strip() or "Cần người dùng bổ sung dàn ý đã duyệt."

    def _citation_lines(self, citations: list[dict[str, Any]]) -> list[str]:
        lines: list[str] = []
        for index, citation in enumerate(citations, start=1):
            title = citation.get("title") or "Nguồn chưa có tên"
            number = citation.get("document_number") or "không số"
            section = citation.get("section") or "không rõ mục"
            score = float(citation.get("score") or 0)
            lines.append(f"[{index}] {title} ({number}), mục {section}, score {score:.4f}.")
        return lines[:10]

    @staticmethod
    def _outline_sections(document_type: str) -> list[str]:
        mapping = {
            "Cong van": ["Căn cứ và bối cảnh", "Nội dung đề nghị/triển khai", "Tổ chức thực hiện", "Nơi nhận"],
            "Bao cao": ["Tình hình chung", "Kết quả thực hiện", "Khó khăn hạn chế", "Kiến nghị đề xuất"],
            "Ke hoach": ["Mục đích yêu cầu", "Nội dung nhiệm vụ", "Tiến độ thực hiện", "Tổ chức thực hiện"],
            "To trinh": ["Sự cần thiết", "Nội dung trình", "Căn cứ và tác động", "Kiến nghị phê duyệt"],
            "Thong bao": ["Nội dung thông báo", "Đối tượng thực hiện", "Thời gian hiệu lực", "Tổ chức thực hiện"],
            "Giay moi": ["Thành phần mời", "Thời gian địa điểm", "Nội dung cuộc họp", "Thông tin liên hệ"],
        }
        return mapping.get(document_type, mapping["Cong van"])

    def _citation_from_library_document(self, document: dict[str, Any], *, score: float, reason: str) -> AIDraftCitation:
        return AIDraftCitation(
            source_document_id=int(document["id"]) if document.get("id") else None,
            title=document.get("title") or document.get("file_name"),
            document_number=document.get("document_number"),
            issued_date=document.get("issued_date"),
            file_path=document.get("file_path"),
            checksum=document.get("checksum"),
            score=round(min(score, 1.0), 6),
            reason=reason,
        )

    def _citation_from_search_result(self, result: dict[str, Any]) -> AIDraftCitation:
        document = result.get("document") or {}
        chunk = result.get("chunk") or {}
        citation = result.get("citation") or {}
        checksum = citation.get("checksum") or document.get("checksum")
        library_document = self.library_repository.get_by_checksum(str(checksum)) if checksum else None
        return AIDraftCitation(
            source_document_id=int(library_document["id"]) if library_document else None,
            source_chunk_id=int(chunk["id"]) if chunk.get("id") else None,
            title=citation.get("title") or document.get("title"),
            document_number=citation.get("document_number") or document.get("document_number"),
            issued_date=citation.get("issued_date") or document.get("issued_date"),
            page=citation.get("page") or chunk.get("page"),
            section=citation.get("section") or chunk.get("section"),
            file_path=citation.get("file_path") or document.get("source_path"),
            checksum=checksum,
            score=float(result.get("score") or citation.get("relevance") or 0),
            quote_text=chunk.get("text"),
            reason="Semantic search evidence",
        )

    def _citation_from_row(self, row: dict[str, Any]) -> AIDraftCitation:
        return AIDraftCitation(
            source_document_id=row.get("source_document_id"),
            source_chunk_id=row.get("source_chunk_id"),
            title=row.get("title"),
            document_number=row.get("document_number"),
            issued_date=row.get("issued_date"),
            page=row.get("page"),
            section=row.get("section"),
            file_path=row.get("file_path"),
            checksum=row.get("checksum"),
            score=float(row.get("score") or 0),
            quote_text=row.get("quote_text"),
            reason=row.get("reason"),
        )

    def _require_request(self, request_id: int) -> dict[str, Any]:
        request = self.repository.get_request(request_id)
        if not request:
            raise ValueError("Không tìm thấy yêu cầu AI Draft")
        return request

    def _require_result(self, result_id: int) -> dict[str, Any]:
        result = self.repository.get_result(result_id)
        if not result:
            raise ValueError("Không tìm thấy dự thảo AI")
        return result

    @staticmethod
    def _selected_type(request: dict[str, Any]) -> str:
        document_type = request.get("selected_document_type") or request.get("detected_document_type")
        if not document_type:
            raise ValueError("Cần người dùng chọn loại văn bản")
        return str(document_type)

    @staticmethod
    def _same_type(value: Any, expected: str) -> bool:
        return AIDraftService._normalize(str(value or "")) == AIDraftService._normalize(expected)

    @staticmethod
    def _is_internal_agency(value: Any) -> bool:
        normalized = AIDraftService._normalize(str(value or ""))
        return "ban xay dung dang" in normalized or ("ban" in normalized and "dang" in normalized)

    @staticmethod
    def _tokens(value: str) -> list[str]:
        return re.findall(r"[0-9a-zA-Z]+", AIDraftService._normalize(value))

    @staticmethod
    def _normalize(value: str) -> str:
        text = unicodedata.normalize("NFD", value.lower())
        text = "".join(ch for ch in text if unicodedata.category(ch) != "Mn")
        text = text.replace("đ", "d")
        return re.sub(r"\s+", " ", text).strip()

    @staticmethod
    def _slugify(value: str) -> str:
        slug = re.sub(r"[^0-9a-zA-Z]+", "-", AIDraftService._normalize(value)).strip("-")
        return slug[:90] or "du-thao"
