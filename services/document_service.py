"""Service nghiệp vụ cho phân hệ Soạn thảo văn bản."""

from __future__ import annotations

from pathlib import Path
from typing import BinaryIO, Any
import json
import logging
import re

from core.config import DOCUMENT_EXPORT_DIR, DOCUMENT_TEMPLATE_DIR
from models.document_draft import DocumentDraft
from models.document_section import DocumentSection
from models.document_template import ALLOWED_TEMPLATE_EXTENSIONS, DocumentTemplate
from services.document_upload_service import DocumentUploadService
from repositories.document_repository import DocumentRepository

logger = logging.getLogger(__name__)


DEFAULT_DOCUMENT_SETTINGS: dict[str, str] = {
    "agency_name": "BAN XÂY DỰNG ĐẢNG",
    "document_code_prefix": "BXD",
    "location_name": "Vĩnh Hòa",
    "recipient_placeholder": "Các cơ quan, đơn vị liên quan",
    "signer_name": "Nguyễn Trung Hiền",
    "signer_title": "TRƯỞNG BAN",
    "font_name": "Times New Roman",
    "font_size": "13",
    "margin_top_cm": "2.0",
    "margin_bottom_cm": "2.0",
    "margin_left_cm": "3.0",
    "margin_right_cm": "2.0",
}


class DocumentService:
    """Điều phối nghiệp vụ mẫu văn bản và dự thảo."""

    def __init__(self, repository: DocumentRepository | None = None) -> None:
        self.repository = repository or DocumentRepository()

    def upload_template_file(
        self,
        source_path: str | Path,
        *,
        document_type: str = "Công văn",
        field: str | None = None,
        created_by: str = "Nguyễn Trung Hiền",
    ) -> int:
        """Upload file mẫu vào kho templates rồi ghi nhận vào database."""
        upload_result = DocumentUploadService().save_local_file(source_path)
        return self.register_template(
            name=Path(upload_result.original_name).stem,
            source_path=upload_result.stored_path,
            document_type=document_type,
            field=field,
            created_by=created_by,
        )

    def upload_template_stream(
        self,
        file_obj: BinaryIO,
        original_name: str,
        *,
        document_type: str = "Công văn",
        field: str | None = None,
        created_by: str = "Nguyễn Trung Hiền",
    ) -> int:
        """Nhận file từ giao diện, lưu vào kho mẫu và đăng ký vào database."""
        upload_result = DocumentUploadService().save_uploaded_file(file_obj, original_name)
        return self.register_template(
            name=Path(upload_result.original_name).stem,
            source_path=upload_result.stored_path,
            document_type=document_type,
            field=field,
            created_by=created_by,
        )

    def register_template(
        self,
        *,
        name: str,
        source_path: str,
        document_type: str = "Công văn",
        field: str | None = None,
        created_by: str = "Nguyễn Trung Hiền",
    ) -> int:
        """Ghi nhận một mẫu văn bản đã có trong thư mục tài liệu."""
        path = Path(source_path)
        ext = path.suffix.lower()
        if ext not in ALLOWED_TEMPLATE_EXTENSIONS:
            logger.error("Không thể đăng ký mẫu do định dạng không hỗ trợ: %s", source_path)
            raise ValueError("Chỉ hỗ trợ DOC, DOCX và PDF")
        file_size = path.stat().st_size if path.exists() else 0
        template = DocumentTemplate(
            name=name or path.stem,
            document_type=document_type,
            field=field,
            source_path=str(path),
            file_ext=ext,
            file_size=file_size,
            created_by=created_by,
        )
        if not template.is_ready_for_analysis:
            raise ValueError("Mẫu văn bản chưa đủ thông tin để phân tích")
        return self.repository.create_template(template)

    # TODO REMOVE AFTER BUILD 1.0: Currently no source path calls this direct model wrapper.
    def create_template_record(self, template: DocumentTemplate) -> int:
        """Tạo mẫu văn bản từ model đã được chuẩn hóa."""
        if not template.is_ready_for_analysis:
            raise ValueError("Mẫu văn bản chưa đủ thông tin")
        return self.repository.create_template(template)

    def list_templates(self, *, status: str | None = None, document_type: str | None = None) -> list[dict]:
        """Danh sách mẫu văn bản."""
        return self.repository.list_templates(status=status, document_type=document_type)

    def get_template(self, template_id: int) -> dict | None:
        """Chi tiết mẫu văn bản."""
        return self.repository.get_template(template_id)

    def mark_template_analyzed(
        self,
        template_id: int,
        *,
        summary: str | None = None,
        keywords: str | None = None,
        analysis_json: str | None = None,
    ) -> int:
        """Đánh dấu mẫu đã phân tích ở các patch Document Intelligence sau."""
        return self.repository.update_template(
            template_id,
            {
                "status": "Đã phân tích",
                "summary": summary,
                "keywords": keywords,
                "analysis_json": analysis_json,
            },
        )

    def analyze_template(self, template_id: int) -> dict[str, Any]:
        """Phân tích mẫu văn bản ở mức offline để người dùng duyệt trước khi dùng."""
        template = self.get_template(template_id)
        if not template:
            raise ValueError("Không tìm thấy mẫu văn bản")
        source_path = Path(str(template["source_path"]))
        raw_text = self._read_template_text(source_path)
        sections = self._extract_sections(raw_text)
        if not sections:
            sections = [
                DocumentSection(
                    section_type="Nội dung",
                    title="Nội dung mẫu",
                    content=raw_text[:2000] if raw_text else "Chưa đọc được nội dung mẫu",
                )
            ]
        self.add_sections(template_id, sections)
        keywords = self._extract_keywords(raw_text, str(template.get("name") or ""))
        analysis = {
            "source_path": str(source_path),
            "document_type": template.get("document_type"),
            "field": template.get("field"),
            "sections": [section.to_dict() for section in sections],
            "format_signals": self._detect_format_signals(raw_text),
            "keywords": keywords,
        }
        summary = self._build_template_summary(raw_text, sections)
        self.mark_template_analyzed(
            template_id,
            summary=summary,
            keywords=", ".join(keywords),
            analysis_json=json.dumps(analysis, ensure_ascii=False),
        )
        logger.info("Analyzed document template id=%s sections=%s", template_id, len(sections))
        return analysis

    def add_sections(self, template_id: int, sections: list[DocumentSection]) -> None:
        """Lưu bố cục/section của mẫu văn bản."""
        self.repository.replace_sections(template_id, sections)

    def list_sections(self, template_id: int) -> list[dict]:
        """Danh sách section của mẫu."""
        return self.repository.list_sections(template_id)

    def create_draft(
        self,
        *,
        title: str,
        document_type: str = "Công văn",
        template_id: int | None = None,
        workspace_id: int | None = None,
        request_text: str | None = None,
        created_by: str = "Nguyễn Trung Hiền",
    ) -> int:
        """Tạo bản ghi dự thảo để người dùng/AI tiếp tục hoàn thiện."""
        draft = DocumentDraft(
            title=title,
            document_type=document_type,
            template_id=template_id,
            workspace_id=workspace_id,
            request_text=request_text,
            created_by=created_by,
        )
        if not draft.title:
            raise ValueError("Tiêu đề dự thảo là bắt buộc")
        return self.repository.create_draft(draft)

    def update_draft_content(self, draft_id: int, content: str, *, status: str = "Chờ duyệt") -> int:
        """Cập nhật nội dung dự thảo do người dùng hoặc AI tham mưu sinh ra."""
        if not content.strip():
            raise ValueError("Nội dung dự thảo không được rỗng")
        return self.repository.update_draft(draft_id, {"draft_content": content, "status": status})

    def get_document_settings(self) -> dict[str, str]:
        """Lấy cấu hình thể thức văn bản, tự bù giá trị mặc định khi chưa cấu hình."""
        settings = {**DEFAULT_DOCUMENT_SETTINGS, **self.repository.list_settings()}
        return {key: str(value).strip() for key, value in settings.items()}

    def update_document_settings(self, settings: dict[str, str]) -> dict[str, str]:
        """Cập nhật cấu hình thể thức văn bản do người dùng duyệt."""
        allowed_keys = set(DEFAULT_DOCUMENT_SETTINGS)
        cleaned: dict[str, str] = {}
        for key, value in settings.items():
            if key not in allowed_keys:
                logger.warning("Ignore unsupported document setting: %s", key)
                continue
            text_value = str(value).strip()
            if not text_value:
                raise ValueError(f"Cấu hình {key} không được để trống")
            if key.startswith("margin_") or key == "font_size":
                number = float(text_value)
                if number <= 0:
                    raise ValueError(f"Cấu hình {key} phải lớn hơn 0")
                text_value = str(number).rstrip("0").rstrip(".")
            self.repository.upsert_setting(key, text_value)
            cleaned[key] = text_value
        return {**self.get_document_settings(), **cleaned}

    def suggest_templates(
        self,
        *,
        request_text: str,
        document_type: str | None = None,
        field: str | None = None,
        limit: int = 5,
    ) -> list[dict[str, Any]]:
        """Gợi ý mẫu gần nhất dựa trên loại văn bản, lĩnh vực và từ khóa yêu cầu."""
        if limit <= 0:
            raise ValueError("limit phải lớn hơn 0")
        request_words = set(self._tokenize(request_text))
        scored: list[dict[str, Any]] = []
        for template in self.list_templates(document_type=document_type):
            score = 0
            reasons: list[str] = []
            if document_type and template.get("document_type") == document_type:
                score += 35
                reasons.append("Cùng loại văn bản")
            if field and template.get("field") == field:
                score += 25
                reasons.append("Cùng lĩnh vực")
            text = " ".join(
                str(template.get(key) or "")
                for key in ("name", "summary", "keywords", "document_type", "field")
            )
            matched = request_words.intersection(self._tokenize(text))
            if matched:
                score += min(len(matched) * 8, 32)
                reasons.append("Trùng từ khóa: " + ", ".join(sorted(matched)[:5]))
            if template.get("status") == "Đã phân tích":
                score += 8
                reasons.append("Đã phân tích bố cục")
            scored.append(
                {
                    **template,
                    "match_score": min(score, 100),
                    "match_reason": "; ".join(reasons) if reasons else "Mẫu sẵn có trong kho",
                }
            )
        return sorted(scored, key=lambda item: (item["match_score"], item.get("created_at") or ""), reverse=True)[:limit]

    def missing_information(self, *, request_text: str, document_type: str) -> list[str]:
        """Liệt kê thông tin còn thiếu để người dùng bổ sung trước khi sinh dự thảo."""
        text = request_text.lower()
        required = {
            "Công văn": {
                "đơn vị nhận/nơi nhận": ["kính gửi", "gửi", "nơi nhận"],
                "nội dung đề nghị": ["đề nghị", "phối hợp", "triển khai", "báo cáo"],
                "thời hạn xử lý": ["trước ngày", "hạn", "thời hạn", "ngày"],
            },
            "Báo cáo": {
                "phạm vi báo cáo": ["tình hình", "kết quả", "đánh giá"],
                "số liệu hoặc nội dung chính": ["số liệu", "kết quả", "nội dung"],
                "kiến nghị/đề xuất": ["kiến nghị", "đề xuất", "phương hướng"],
            },
            "Kế hoạch": {
                "mục đích/yêu cầu": ["mục đích", "yêu cầu"],
                "nội dung thực hiện": ["nội dung", "nhiệm vụ"],
                "thời gian thực hiện": ["thời gian", "tiến độ", "ngày"],
            },
        }
        rules = required.get(document_type, required["Công văn"])
        return [label for label, signals in rules.items() if not any(signal in text for signal in signals)]

    def generate_draft(
        self,
        *,
        title: str,
        document_type: str,
        request_text: str,
        template_id: int | None = None,
        workspace_id: int | None = None,
        extra_info: str | None = None,
        created_by: str = "Nguyễn Trung Hiền",
    ) -> int:
        """Sinh dự thảo tham mưu từ yêu cầu và mẫu đã chọn, trạng thái luôn chờ duyệt."""
        if not request_text.strip():
            raise ValueError("Yêu cầu soạn thảo không được rỗng")
        content = self._compose_draft_content(
            title=title,
            document_type=document_type,
            request_text=request_text,
            template_id=template_id,
            extra_info=extra_info,
            settings=self.get_document_settings(),
        )
        draft_id = self.create_draft(
            title=title,
            document_type=document_type,
            template_id=template_id,
            workspace_id=workspace_id,
            request_text=request_text,
            created_by=created_by,
        )
        self.update_draft_content(draft_id, content, status="Chờ duyệt")
        logger.info("Generated document draft id=%s template_id=%s", draft_id, template_id)
        return draft_id

    def get_draft(self, draft_id: int) -> dict[str, Any] | None:
        """Chi tiết dự thảo."""
        return self.repository.get_draft(draft_id)

    def review_draft_format(self, draft_id: int) -> list[dict[str, str]]:
        """Kiểm tra nhanh thể thức dự thảo trước khi xuất file."""
        draft = self.get_draft(draft_id)
        if not draft:
            raise ValueError("Không tìm thấy dự thảo")
        content = str(draft.get("draft_content") or "")
        settings = self.get_document_settings()
        checks = [
            ("Quốc hiệu", "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", "Thiếu quốc hiệu"),
            ("Tiêu ngữ", "Độc lập - Tự do - Hạnh phúc", "Thiếu tiêu ngữ"),
            ("Số ký hiệu", "Số:", "Thiếu dòng số/ký hiệu văn bản"),
            ("Nơi nhận", "Nơi nhận:", "Thiếu mục nơi nhận"),
            ("Chữ ký", settings["signer_title"], "Thiếu khối chữ ký/người ký"),
        ]
        result: list[dict[str, str]] = []
        for name, signal, warning in checks:
            ok = signal.lower() in content.lower()
            result.append({"name": name, "status": "Đạt" if ok else "Cần bổ sung", "note": "" if ok else warning})
        if len(content.strip()) < 200:
            result.append({"name": "Dung lượng nội dung", "status": "Cần bổ sung", "note": "Nội dung dự thảo còn ngắn"})
        else:
            result.append({"name": "Dung lượng nội dung", "status": "Đạt", "note": ""})
        return result

    def export_draft_docx(self, draft_id: int) -> Path:
        """Xuất dự thảo sang DOCX và lưu đường dẫn vào bản ghi dự thảo."""
        draft = self.get_draft(draft_id)
        if not draft:
            raise ValueError("Không tìm thấy dự thảo")
        content = str(draft.get("draft_content") or "").strip()
        if not content:
            raise ValueError("Dự thảo chưa có nội dung để xuất")
        try:
            from docx import Document
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.shared import Cm
            from docx.shared import Pt
        except ImportError as exc:
            raise RuntimeError("Thiếu thư viện python-docx để xuất DOCX") from exc

        settings = self.get_document_settings()
        DOCUMENT_EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        safe_name = self._slugify(str(draft.get("title") or "du-thao")) or "du-thao"
        output = DOCUMENT_EXPORT_DIR / f"{draft_id:04d}_{safe_name}.docx"
        document = Document()
        section = document.sections[0]
        section.top_margin = Cm(float(settings["margin_top_cm"]))
        section.bottom_margin = Cm(float(settings["margin_bottom_cm"]))
        section.left_margin = Cm(float(settings["margin_left_cm"]))
        section.right_margin = Cm(float(settings["margin_right_cm"]))
        style = document.styles["Normal"]
        style.font.name = settings["font_name"]
        style.font.size = Pt(float(settings["font_size"]))
        for paragraph in content.splitlines():
            paragraph = paragraph.strip()
            if not paragraph:
                document.add_paragraph()
                continue
            doc_paragraph = document.add_paragraph()
            run = doc_paragraph.add_run(paragraph)
            if paragraph.isupper():
                run.bold = True
                doc_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif paragraph.startswith(("CỘNG HÒA", "Độc lập")):
                run.bold = "CỘNG HÒA" in paragraph
                doc_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            elif paragraph.startswith((settings["location_name"], "Số:", "Kính gửi")):
                run.bold = paragraph.startswith("Kính gửi")
        document.save(output)
        self.repository.update_draft(draft_id, {"output_path": str(output), "status": "Đã xuất"})
        logger.info("Exported document draft id=%s to %s", draft_id, output)
        return output

    def list_drafts(self, *, status: str | None = None) -> list[dict]:
        """Danh sách dự thảo."""
        return self.repository.list_drafts(status=status)

    # TODO REMOVE AFTER BUILD 1.0: Runtime directories are already created by core.config.
    def ensure_template_directory(self) -> Path:
        """Đảm bảo thư mục lưu mẫu tồn tại."""
        DOCUMENT_TEMPLATE_DIR.mkdir(parents=True, exist_ok=True)
        return DOCUMENT_TEMPLATE_DIR

    def _read_template_text(self, path: Path) -> str:
        """Đọc nội dung mẫu DOCX/PDF ở mức đủ để phân tích bố cục."""
        if not path.exists():
            raise FileNotFoundError(f"Không tìm thấy file mẫu: {path}")
        if path.suffix.lower() == ".docx":
            try:
                from docx import Document
            except ImportError as exc:
                raise RuntimeError("Thiếu thư viện python-docx để đọc DOCX") from exc
            document = Document(path)
            texts = [p.text.strip() for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        texts.append(" | ".join(cells))
            return "\n".join(texts)
        if path.suffix.lower() == ".pdf":
            try:
                import fitz
            except ImportError as exc:
                raise RuntimeError("Thiếu thư viện PyMuPDF để đọc PDF") from exc
            with fitz.open(path) as document:
                return "\n".join(page.get_text("text") for page in document)
        logger.warning("File DOC chưa hỗ trợ đọc sâu, chỉ phân tích metadata: %s", path)
        return path.stem

    def _extract_sections(self, text: str) -> list[DocumentSection]:
        """Tách các phần thường gặp trong văn bản hành chính."""
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        sections: list[DocumentSection] = []
        current_title = "Mở đầu"
        current_type = "Mở đầu"
        buffer: list[str] = []
        heading_patterns = [
            ("Căn cứ", "Căn cứ"),
            ("Kính gửi", "Nơi nhận xử lý"),
            ("Nội dung", "Nội dung"),
            ("Tổ chức thực hiện", "Tổ chức thực hiện"),
            ("Nơi nhận", "Nơi nhận"),
        ]
        for line in lines:
            matched = next((item for item in heading_patterns if item[0].lower() in line.lower()), None)
            is_heading = matched is not None or (line.isupper() and len(line) <= 120)
            if is_heading and buffer:
                sections.append(
                    DocumentSection(section_order=len(sections) + 1, section_type=current_type, title=current_title, content="\n".join(buffer))
                )
                buffer = []
            if is_heading:
                current_title = line[:120]
                current_type = matched[1] if matched else "Tiêu đề"
            else:
                buffer.append(line)
        if buffer or not sections:
            sections.append(
                DocumentSection(section_order=len(sections) + 1, section_type=current_type, title=current_title, content="\n".join(buffer))
            )
        return sections[:20]

    def _detect_format_signals(self, text: str) -> dict[str, bool]:
        lower_text = text.lower()
        return {
            "has_national_header": "cộng hòa xã hội chủ nghĩa việt nam" in lower_text,
            "has_motto": "độc lập" in lower_text and "tự do" in lower_text and "hạnh phúc" in lower_text,
            "has_document_number": "số:" in lower_text or "số " in lower_text,
            "has_recipient": "kính gửi" in lower_text or "nơi nhận" in lower_text,
            "has_signature_block": "người ký" in lower_text or "trưởng ban" in lower_text or "kt." in lower_text,
        }

    def _build_template_summary(self, text: str, sections: list[DocumentSection]) -> str:
        first_lines = [line.strip() for line in text.splitlines() if line.strip()][:5]
        section_names = ", ".join(section.section_type for section in sections[:6])
        if first_lines:
            return f"Bố cục nhận diện: {section_names}. Mở đầu mẫu: {' / '.join(first_lines)[:300]}"
        return f"Bố cục nhận diện: {section_names}."

    def _compose_draft_content(
        self,
        *,
        title: str,
        document_type: str,
        request_text: str,
        template_id: int | None,
        extra_info: str | None,
        settings: dict[str, str],
    ) -> str:
        template = self.get_template(template_id) if template_id else None
        template_note = f"Mẫu tham chiếu: {template['name']}" if template else "Mẫu tham chiếu: chưa chọn"
        sections = self.list_sections(template_id) if template_id else []
        section_note = ", ".join(str(section.get("section_type")) for section in sections[:6]) or "Mở đầu, nội dung, nơi nhận"
        missing = self.missing_information(request_text=request_text, document_type=document_type)
        missing_note = "Thông tin cần người dùng kiểm tra thêm: " + "; ".join(missing) if missing else "Thông tin đầu vào đã đủ để lập dự thảo sơ bộ."
        body = [
            "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM",
            "Độc lập - Tự do - Hạnh phúc",
            "",
            f"{settings['agency_name']}",
            f"{settings['location_name']}, ngày ...... tháng ...... năm ......",
            f"Số: ......../{settings['document_code_prefix']}",
            "",
            title.upper(),
            "",
            f"Kính gửi: {settings['recipient_placeholder']}",
            "",
            f"Căn cứ yêu cầu soạn thảo: {request_text.strip()}",
            "",
            f"Hệ thống tham mưu dự thảo {document_type.lower()} theo bố cục: {section_note}.",
            template_note,
        ]
        if extra_info and extra_info.strip():
            body.extend(["", "Thông tin bổ sung:", extra_info.strip()])
        body.extend(
            [
                "",
                "Nội dung dự thảo:",
                "1. Tình hình, căn cứ và yêu cầu xử lý",
                "Trình bày rõ bối cảnh, căn cứ triển khai và nội dung cần thực hiện theo chỉ đạo.",
                "",
                "2. Nội dung đề nghị/triển khai",
                "Các đơn vị, cá nhân liên quan căn cứ chức năng, nhiệm vụ được giao để phối hợp thực hiện; bảo đảm tiến độ, chất lượng và có minh chứng kèm theo.",
                "",
                "3. Tổ chức thực hiện",
                "Giao bộ phận phụ trách theo dõi, tổng hợp kết quả, báo cáo lãnh đạo xem xét trước khi ban hành chính thức.",
                "",
                missing_note,
                "",
                "Nơi nhận:",
                "- Như trên;",
                "- Lưu: VT.",
                "",
                settings["signer_title"],
                settings["signer_name"],
            ]
        )
        return "\n".join(body)

    @staticmethod
    def _extract_keywords(*values: str) -> list[str]:
        words: list[str] = []
        for value in values:
            words.extend(DocumentService._tokenize(value))
        ignored = {"văn", "bản", "công", "việc", "theo", "cho", "các", "mẫu", "nội", "dung"}
        unique = []
        for word in words:
            if len(word) >= 3 and word not in ignored and word not in unique:
                unique.append(word)
        return unique[:12]

    @staticmethod
    def _tokenize(value: str) -> list[str]:
        return re.findall(r"[0-9A-Za-zÀ-ỹ]+", value.lower())

    @staticmethod
    def _slugify(value: str) -> str:
        value = value.strip().lower()
        value = re.sub(r"[^0-9a-zA-ZÀ-ỹ]+", "-", value, flags=re.UNICODE)
        return re.sub(r"-+", "-", value).strip("-")[:90]
