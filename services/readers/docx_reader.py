"""Reader DOCX cho kho văn bản."""

from __future__ import annotations

from pathlib import Path
import logging

from services.readers.base import ReaderResult

logger = logging.getLogger(__name__)


class DocxReader:
    """Đọc nội dung DOCX bằng python-docx."""

    def read(self, path: str | Path) -> ReaderResult:
        file_path = Path(path)
        if file_path.suffix.lower() == ".doc":
            return ReaderResult(text="", status="unsupported", error="DOC cũ chưa hỗ trợ đọc nội dung")
        try:
            from docx import Document

            document = Document(file_path)
            texts = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
            tables: list[list[list[str]]] = []
            for table in document.tables:
                table_rows: list[list[str]] = []
                for row in table.rows:
                    values = [cell.text.strip() for cell in row.cells]
                    table_rows.append(values)
                    row_text = " | ".join(value for value in values if value)
                    if row_text:
                        texts.append(row_text)
                tables.append(table_rows)
            return ReaderResult(
                text="\n".join(texts),
                pages=1,
                tables=tables,
                metadata={"paragraphs": len(document.paragraphs), "tables": len(document.tables)},
                status="read",
            )
        except Exception as exc:
            logger.exception("Cannot read DOCX file: %s", file_path)
            return ReaderResult(text="", status="failed", error=str(exc))
