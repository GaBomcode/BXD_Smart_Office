"""Reader XLSX cho kho văn bản."""

from __future__ import annotations

from pathlib import Path
import logging

from services.readers.base import ReaderResult

logger = logging.getLogger(__name__)


class XlsxReader:
    """Đọc XLSX bằng openpyxl; XLS cũ tạm đánh dấu unsupported."""

    def read(self, path: str | Path) -> ReaderResult:
        file_path = Path(path)
        if file_path.suffix.lower() == ".xls":
            return ReaderResult(text="", status="unsupported", error="XLS cũ chưa hỗ trợ đọc nội dung")
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(file_path, read_only=True, data_only=True)
            texts: list[str] = []
            tables: list[list[list[str]]] = []
            for sheet in workbook.worksheets:
                rows: list[list[str]] = []
                for row in sheet.iter_rows(values_only=True):
                    values = ["" if value is None else str(value) for value in row]
                    if any(value.strip() for value in values):
                        rows.append(values)
                        texts.append(" | ".join(value for value in values if value.strip()))
                tables.append(rows)
            return ReaderResult(
                text="\n".join(texts),
                pages=len(workbook.worksheets),
                tables=tables,
                metadata={"sheets": workbook.sheetnames},
                status="read",
            )
        except Exception as exc:
            logger.exception("Cannot read XLSX file: %s", file_path)
            return ReaderResult(text="", status="failed", error=str(exc))
