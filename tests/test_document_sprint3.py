"""Kiểm thử luồng Sprint 3 - Soạn thảo văn bản."""

from pathlib import Path

from docx import Document

from database.init_db import init_database
from services.document_service import DocumentService


def _create_sample_docx(path: Path) -> None:
    document = Document()
    document.add_paragraph("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
    document.add_paragraph("Độc lập - Tự do - Hạnh phúc")
    document.add_paragraph("Số: 01-CV/BXD")
    document.add_paragraph("CÔNG VĂN")
    document.add_paragraph("Kính gửi: Các đơn vị liên quan")
    document.add_paragraph("Nội dung triển khai nhiệm vụ công tác tổ chức.")
    document.add_paragraph("Nơi nhận:")
    document.add_paragraph("TRƯỞNG BAN")
    document.save(path)


def test_analyze_template_saves_sections_and_keywords(tmp_path: Path) -> None:
    """Phân tích mẫu phải lưu section, summary và trạng thái đã phân tích."""
    init_database()
    sample = tmp_path / "mau_cong_van.docx"
    _create_sample_docx(sample)

    service = DocumentService()
    template_id = service.register_template(
        name="Mẫu công văn tổ chức",
        source_path=str(sample),
        document_type="Công văn",
        field="Tổ chức",
    )
    analysis = service.analyze_template(template_id)
    template = service.get_template(template_id)
    sections = service.list_sections(template_id)

    assert template is not None
    assert template["status"] == "Đã phân tích"
    assert template["summary"]
    assert analysis["format_signals"]["has_national_header"] is True
    assert len(sections) >= 1


def test_suggest_generate_review_and_export_draft(tmp_path: Path) -> None:
    """Service phải gợi ý mẫu, sinh dự thảo chờ duyệt, kiểm tra và xuất DOCX."""
    init_database()
    sample = tmp_path / "mau_cong_van.docx"
    _create_sample_docx(sample)

    service = DocumentService()
    template_id = service.register_template(
        name="Mẫu công văn tổ chức",
        source_path=str(sample),
        document_type="Công văn",
        field="Tổ chức",
    )
    service.analyze_template(template_id)

    suggestions = service.suggest_templates(
        request_text="Soạn công văn đề nghị phối hợp triển khai nhiệm vụ tổ chức trước ngày 30",
        document_type="Công văn",
        field="Tổ chức",
    )
    assert suggestions
    assert suggestions[0]["id"] == template_id

    draft_id = service.generate_draft(
        title="Công văn phối hợp triển khai nhiệm vụ",
        document_type="Công văn",
        request_text="Soạn công văn đề nghị phối hợp triển khai nhiệm vụ tổ chức trước ngày 30",
        template_id=template_id,
    )
    draft = service.get_draft(draft_id)
    assert draft is not None
    assert draft["status"] == "Chờ duyệt"
    assert "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM" in draft["draft_content"]

    checks = service.review_draft_format(draft_id)
    assert all(item["status"] == "Đạt" for item in checks if item["name"] in {"Quốc hiệu", "Tiêu ngữ"})

    output = service.export_draft_docx(draft_id)
    assert output.exists()
    assert service.get_draft(draft_id)["status"] == "Đã xuất"  # type: ignore[index]


def test_document_settings_are_used_when_generating_and_exporting_draft(tmp_path: Path) -> None:
    """Cấu hình thể thức phải được dùng khi sinh dự thảo và xuất DOCX."""
    init_database()
    service = DocumentService()
    service.update_document_settings(
        {
            "agency_name": "BAN XÂY DỰNG ĐẢNG XÃ VĨNH HÒA",
            "document_code_prefix": "BXD-VH",
            "location_name": "Vĩnh Hòa",
            "recipient_placeholder": "Chi bộ trực thuộc",
            "signer_title": "TRƯỞNG BAN",
            "signer_name": "Nguyễn Trung Hiền",
            "font_name": "Times New Roman",
            "font_size": "13",
            "margin_top_cm": "2",
            "margin_bottom_cm": "2",
            "margin_left_cm": "3",
            "margin_right_cm": "2",
        }
    )

    draft_id = service.generate_draft(
        title="Công văn kiểm tra cấu hình",
        document_type="Công văn",
        request_text="Soạn công văn đề nghị phối hợp báo cáo trước ngày 30 gửi Chi bộ trực thuộc",
    )
    draft = service.get_draft(draft_id)
    assert draft is not None
    assert "BAN XÂY DỰNG ĐẢNG XÃ VĨNH HÒA" in draft["draft_content"]
    assert "Số: ......../BXD-VH" in draft["draft_content"]
    assert "Kính gửi: Chi bộ trực thuộc" in draft["draft_content"]

    output = service.export_draft_docx(draft_id)
    exported = Document(output)
    exported_text = "\n".join(paragraph.text for paragraph in exported.paragraphs)
    assert "Nguyễn Trung Hiền" in exported_text
    assert "BXD-VH" in exported_text
