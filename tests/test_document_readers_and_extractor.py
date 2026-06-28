"""Kiểm thử reader và metadata extractor Sprint 4."""

from pathlib import Path

from docx import Document
from openpyxl import Workbook

from services.document_metadata_extractor import DocumentMetadataExtractor
from services.readers import DocxReader, TxtReader, XlsxReader


def test_docx_reader_returns_normalized_result(tmp_path: Path) -> None:
    path = tmp_path / "mau.docx"
    document = Document()
    document.add_paragraph("BAN XÂY DỰNG ĐẢNG")
    document.add_paragraph("Số: 12-CV/BXD")
    document.add_paragraph("V/v kiểm thử reader")
    document.save(path)

    result = DocxReader().read(path)

    assert result.status == "read"
    assert "12-CV/BXD" in result.text
    assert result.error is None


def test_xlsx_and_txt_readers(tmp_path: Path) -> None:
    xlsx = tmp_path / "bang.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet["A1"] = "Số"
    sheet["B1"] = "Nội dung"
    workbook.save(xlsx)
    txt = tmp_path / "vanban.txt"
    txt.write_text("Thông báo kiểm thử", encoding="utf-8")

    assert XlsxReader().read(xlsx).status == "read"
    assert "Nội dung" in XlsxReader().read(xlsx).text
    assert TxtReader().read(txt).text == "Thông báo kiểm thử"


def test_metadata_extractor_handles_party_admin_text() -> None:
    text = """
    ĐẢNG ỦY XÃ VĨNH HÒA
    BAN XÂY DỰNG ĐẢNG
    Số: 15-CV/BXD
    Vĩnh Hòa, ngày 12 tháng 5 năm 2026
    CÔNG VĂN
    V/v phối hợp báo cáo công tác tổ chức cán bộ
    TRƯỞNG BAN
    Nguyễn Trung Hiền
    """

    metadata = DocumentMetadataExtractor().extract(text, file_name="15-CV-BXD.docx")

    assert metadata["document_number"] == "15-CV/BXD"
    assert metadata["issued_date"] == "2026-05-12"
    assert metadata["document_type"] == "Công văn"
    assert metadata["field"] == "Tổ chức"
    assert metadata["signer"] == "Nguyễn Trung Hiền"
